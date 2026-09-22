"""
공사기간 적정성 검토 보고서 — DOCX 생성 (python-docx)

입력(JSON): 프런트 reportModel 요약 + 표 데이터
  {
    title, site, meta:{region, station, file, start, plan, verdict, verdict_desc, total, diff},
    narrative:{overview, weather, opinion},          # 템플릿 또는 LLM 서술
    mapping:{rows, auto, review, unmapped, unmapped_names:[...]},
    work_rows:[{name, std, ref, qty, unit, prod, prod_unit, crews, hz, days}],
    work_total:{net, indirect, adjust, total, hz_text},
    weather:{rules:[...], scopes:{...}, monthly:[{m, wx, hol}], annual_wx, annual_hol, peak:[...], source},
    cpm:[{key, pred, dur, es, lf, float, start, end, critical}], cp_days, critical_path:[...],
    totals:[{label, days, basis}], calendar_total, notes:[...]
  }
"""
from io import BytesIO
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = '맑은 고딕'


def _font(run, size=10, bold=False, color=None):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if color: run.font.color.rgb = RGBColor(*color)


def _shade(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill); tcPr.append(shd)


def _p(doc, text, size=10, bold=False, align=None, color=None, space_after=4):
    p = doc.add_paragraph(); r = p.add_run(str(text)); _font(r, size, bold, color)
    if align == 'center': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _h(doc, text):
    p = doc.add_paragraph(); r = p.add_run(text); _font(r, 12, True, (15, 61, 110))
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    pPr = p._element.get_or_add_pPr(); bdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '8'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), '0F3D6E'); bdr.append(b); pPr.append(bdr)


def _table(doc, header, rows, widths=None, num_cols=()):
    t = doc.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(str(hd)); _font(r, 9, True); _shade(c, 'EEF2F7')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; par = cells[i].paragraphs[0]; r = par.add_run('' if v is None else str(v)); _font(r, 9)
            if i in num_cols: par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _ai_box(doc, label, text):
    t = doc.add_table(rows=1, cols=1); t.style = 'Table Grid'; c = t.rows[0].cells[0]; _shade(c, 'F3F7FD'); c.text = ''
    r = c.paragraphs[0].add_run(label); _font(r, 8, True, (28, 92, 171))
    p = c.add_paragraph(); r = p.add_run(text); _font(r, 10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(d: dict) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2); s.right_margin = Cm(2)
    st = doc.styles['Normal']; st.font.name = FONT; st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT); st.font.size = Pt(10)
    m = d.get('meta', {})
    _p(doc, d.get('title', '공사기간 적정성 검토 보고서'), 18, True, 'center', space_after=2)
    _p(doc, f"{d.get('site','')} · 검토일 {m.get('date','')} · 문서번호 {m.get('docno','')}", 9, False, 'center', (90, 90, 90), 10)
    # 판정 박스
    t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
    c0, c1 = t.rows[0].cells; c0.width = Cm(4); c1.width = Cm(13)
    col = {'적정': (12, 122, 12), '검토 필요': (122, 75, 0)}.get(m.get('verdict'), (176, 36, 36))
    c0.text = ''; r = c0.paragraphs[0].add_run(m.get('verdict', '')); _font(r, 16, True, col); c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c1.text = ''; r = c1.paragraphs[0].add_run(f"산정 공사기간 {m.get('total')}일 vs 계획 공기 {m.get('plan')}일 ({'+' if (m.get('diff') or 0) >= 0 else ''}{m.get('diff')}일)"); _font(r, 10, True)
    r = c1.add_paragraph().add_run(m.get('verdict_desc', '')); _font(r, 9)
    doc.add_paragraph()
    _h(doc, '1. 검토 개요')
    _table(doc, ['항목', '내용', '항목', '내용'], [['현장', d.get('site', ''), '현장 위치', m.get('region', '')], ['검토 대상', m.get('file', ''), '착수 예정일', m.get('start', '')], ['현장 특성', m.get('traits', ''), '기상 지점', m.get('station', '')]], [3, 6, 3, 5])
    n = d.get('narrative', {})
    _ai_box(doc, n.get('label', '서술 초안 · 검토자 확인 필요'), n.get('overview', ''))
    _h(doc, '2. 공종 매핑 결과')
    mp = d.get('mapping', {})
    _p(doc, f"내역서 품명 {mp.get('rows')}건을 건설공사 표준품셈 표준 명칭과 대조한 결과 자동 확정 {mp.get('auto')}건, 검토 필요 {mp.get('review')}건, 미매칭 {mp.get('unmapped')}건입니다.")
    if mp.get('unmapped_names'):
        _p(doc, '미매칭 항목: ' + ', '.join(mp['unmapped_names']) + ' — 표준품셈에 해당 생산성 기준이 없거나 단위가 불일치하여 산정에서 제외되었으며, 정리기간 또는 별도 실적 기준으로 반영이 필요합니다.', 9)
    _h(doc, '3. 공종별 작업일수 산정')
    wr = d.get('work_rows', [])
    _table(doc, ['공종', '적용 표준품셈 (근거)', '수량', '1일 시공량', '투입조', '할증', '순작업일'],
           [[w.get('name'), f"{w.get('std','')}\n{w.get('ref','')}", f"{w.get('qty')} {w.get('unit','')}", f"{w.get('prod','')} {w.get('prod_unit','')}", w.get('crews', ''), w.get('hz', '—'), w.get('days')] for w in wr], [3.2, 6.2, 2.2, 2.2, 1.2, 1.2, 1.6], num_cols=(2, 3, 4, 5, 6))
    wt = d.get('work_total', {})
    _p(doc, f"순작업일수는 표준품셈 1-4절 품 할증({wt.get('hz_text','해당 없음')})을 시공량에 반영한 값이며, 합계 {wt.get('net')}일에 검측·양생 등 간접작업일 {wt.get('indirect')}일, 보정 {wt.get('adjust')}일을 가산하여 공정별 작업일수 합계는 {wt.get('total')}일입니다.")
    _h(doc, '4. 기상 조건 분석 및 비작업일수')
    wx = d.get('weather', {})
    if n.get('weather'): _ai_box(doc, n.get('label', '서술 초안'), n['weather'])
    _p(doc, wx.get('text', ''))
    mon = wx.get('monthly', [])
    if mon:
        _table(doc, ['구분(토공·옥외)'] + [f"{x['m']}월" for x in mon] + ['연간'],
               [['기상 비작업일'] + [f"{x['wx']:.1f}" for x in mon] + [f"{wx.get('annual_wx',0):.1f}"], ['휴일'] + [f"{x['hol']:.1f}" for x in mon] + [f"{wx.get('annual_hol',0):.1f}"]], num_cols=tuple(range(1, 14)))
    _h(doc, '5. 공정 분석 (CPM)')
    _p(doc, f"{len(d.get('cpm', []))}개 공정의 선후행 관계를 기준으로 분석한 결과 주공정선은 {' → '.join(d.get('critical_path', []))}이며, 주공정 작업일수는 {d.get('cp_days')}일입니다. 주공정선에 포함된 공정은 여유일수가 0일로 지연 시 전체 공기가 그대로 연장됩니다.")
    _table(doc, ['공정', '선행 공정', '작업일', 'ES', 'LF', '여유', '착수~완료', '주공정'],
           [[a.get('key'), a.get('pred', ''), a.get('dur'), a.get('es'), a.get('lf'), a.get('float'), f"{a.get('start','')}~{a.get('end','')}", '●' if a.get('critical') else ''] for a in d.get('cpm', [])], [3.6, 3.6, 1.4, 1.2, 1.2, 1.2, 3.2, 1.4], num_cols=(2, 3, 4, 5))
    _h(doc, '6. 공사기간 산정 총괄')
    _table(doc, ['구분', '일수', '산정 근거'], [[x.get('label'), x.get('days'), x.get('basis')] for x in d.get('totals', [])], [4, 2, 11], num_cols=(1,))
    _h(doc, '7. 적정성 검토 의견')
    _ai_box(doc, n.get('label', '서술 초안 · 검토자 확인 필요'), n.get('opinion', ''))
    _h(doc, '8. 산정 근거 및 유의사항')
    for x in d.get('notes', []):
        p = doc.add_paragraph(style='List Bullet'); r = p.add_run(x); _font(r, 9)
    doc.add_paragraph()
    _p(doc, '작성: 공사기간 산정 콘솔 (자동 초안) · 검토자: ____________ (인)', 9, False, None, (90, 90, 90))
    buf = BytesIO(); doc.save(buf); return buf.getvalue()
